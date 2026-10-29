# -*- coding: utf-8 -*-
"""把解析结果生成为 .pptx（仿保诚参考PPT版式：A4竖版、红表头、表格铺满、双行分组表头）。
支持单份与多提取方案组合：基础表(投入/正常/乐观/保证)只出一次，每个提取方案各出一页。

版式规范（依据用户提供的参考PPT）：
- 主色 C00000（红）：表头底、产品名、本币(人民币/港币)列、关键金额数字
- 正文 222222，说明 888888，缴费描述普通部分 333333，标签 0D0D0D
- 表头字号：正常/乐观表 24pt、保证vs非保证 18+14pt、提取表 18pt
- 数据字号：正常/乐观表 18pt、保证/提取表 14pt
- 隔行填充 FAFAFA；保证vs非保证表不隔行
"""
import io
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

RED = RGBColor(0xC0, 0x00, 0x00)      # 主红（表头/产品名/本币列/关键数字）
DARK = RGBColor(0x22, 0x22, 0x22)      # 正文
GRAY = RGBColor(0x88, 0x88, 0x88)      # 说明 / 汇率
SUB = RGBColor(0x33, 0x33, 0x33)       # 缴费描述普通文字
INK0 = RGBColor(0x0D, 0x0D, 0x0D)      # 提取页标题
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xFA, 0xFA, 0xFA)     # 隔行填充


def pick_rows(data_rows):
    last_year = data_rows[-1]["year"] if data_rows else None
    has_gap = any(r.get("is_age_marker") and r["year"] != last_year for r in data_rows)
    if has_gap:
        out = [r for r in data_rows if (r.get("is_age_marker") and r["year"] > 30)
               or (not r.get("is_age_marker") and r["year"] <= 30 and r["year"] % 5 == 0)]
    else:
        out = [r for r in data_rows if r["year"] % 5 == 0]
    last = data_rows[-1] if data_rows else None
    if last and not any(r["year"] == last["year"] for r in out):
        out.append(last)
    seen = set(); res = []
    for r in sorted(out, key=lambda x: x["year"]):
        if r["year"] in seen: continue
        seen.add(r["year"]); res.append(r)
    return res


def _wan(v):
    """美元金额缩写：>=1亿显示亿，否则万；整亿/整万去小数，其余保留1位小数。"""
    if v >= 1e8:
        t = int(v // 1e7 + 1e-9) / 10
        return f"{int(t)} 亿" if t == int(t) else f"{t:.1f} 亿"
    t = int(v // 1000 + 1e-9) / 10
    return f"{int(t)} 万" if t == int(t) else f"{t:.1f} 万"


def _rmb(v, rate):
    """人民币/港币换算，显示规则与美元一致。"""
    total = v * rate
    return _wan(total)


def _cl(data):
    return data.get("curLabel", "人民币")


def _style_cell(cell, text, bold=False, color=DARK, size=14, align=PP_ALIGN.CENTER, fill=None):
    cell.text = ""
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = cell.text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "微软雅黑"
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill


def _style_cell_lines(cell, segs, align=PP_ALIGN.CENTER, fill=None):
    """多行单元格，segs = [(text, bold, color, size), ...]，每段一个段落。"""
    cell.vertical_anchor = MSO_ANCHOR.TOP
    tf = cell.text_frame
    tf.word_wrap = True
    for i, (text, bold, color, size) in enumerate(segs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = "微软雅黑"
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill


def _box(slide, text, l, t, w, size=14, color=DARK, bold=False, align=PP_ALIGN.LEFT, h=0.5):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    p = box.text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "微软雅黑"
    return box


def _richbox(slide, l, t, w, size, segs, align=PP_ALIGN.LEFT, h=0.5):
    """多 run 文本框，segs = [(text, color, bold), ...]。用于同一句内部分红/加粗。"""
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    p = box.text_frame.paragraphs[0]
    p.alignment = align
    for txt, color, bold in segs:
        run = p.add_run()
        run.text = txt
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "微软雅黑"
    return box


_NUMRE = re.compile(r"\d[\d,]*(?:\.\d+)?%?")


def _richbox_red(slide, l, t, w, size, segs, align=PP_ALIGN.LEFT, h=0.5):
    """segs=[(text,color,bold)]；内部把每段中的数字/百分比/带千分位金额标红加粗。"""
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    p = box.text_frame.paragraphs[0]
    p.alignment = align
    for txt, color, bold in segs:
        pos = 0
        for mm in _NUMRE.finditer(txt):
            if mm.start() > pos:
                r = p.add_run(); r.text = txt[pos:mm.start()]
                r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = "微软雅黑"
            r = p.add_run(); r.text = mm.group()
            r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = RED; r.font.name = "微软雅黑"
            pos = mm.end()
        if pos < len(txt):
            r = p.add_run(); r.text = txt[pos:]
            r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = "微软雅黑"
    return box


def _add_line(slide, l, t, w, h=0.02, color=RED):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()


def _footer(slide, rate=6.8, cl="人民币"):
    _box(slide, f"美元汇率：{rate}（仅供参考）", 5.1, 11.35, 2.8, size=9, color=GRAY, align=PP_ALIGN.RIGHT)


def _build_table(slide, top_in, headers, data_rows, col_w, row_h=0.44,
                 two_row_header=False, hdr_size=11, hdr2_size=None,
                 data_size=10, zebra=True, red_cols=()):
    """通用表格。red_cols 指定本币列索引（数据用红色），zebra 控制隔行填充。"""
    n_cols = len(headers[0])
    n_rows = len(data_rows) + (2 if two_row_header else 1)
    height = row_h * n_rows
    tbl_shape = slide.shapes.add_table(n_rows, n_cols, Inches(0.52), Inches(top_in),
                                       Inches(7.22), Inches(height))
    tbl = tbl_shape.table
    for i, w in enumerate(col_w):
        tbl.columns[i].width = Inches(w)
    for r in range(n_rows):
        tbl.rows[r].height = Inches(row_h)

    if two_row_header:
        for j, h in enumerate(headers[0]):
            _style_cell(tbl.cell(0, j), h or "", bold=True, color=WHITE, size=hdr_size, fill=RED)
        tbl.cell(0, 0).merge(tbl.cell(1, 0))
        tbl.cell(0, 1).merge(tbl.cell(0, 2))
        tbl.cell(0, 3).merge(tbl.cell(0, 4))
        for j, h in enumerate(headers[1]):
            _style_cell(tbl.cell(1, j), h, bold=True, color=WHITE, size=hdr2_size or hdr_size, fill=RED)
        start = 2
    else:
        for j, h in enumerate(headers[0]):
            _style_cell(tbl.cell(0, j), h, bold=True, color=WHITE, size=hdr_size, fill=RED)
        start = 1

    for i, row in enumerate(data_rows):
        fill = LIGHT if (zebra and i % 2 == 0) else None
        for j, val in enumerate(row):
            color = RED if j in red_cols else DARK
            _style_cell(tbl.cell(start + i, j), val, size=data_size, color=color, fill=fill)
    return tbl_shape


def _new_prs():
    prs = Presentation()
    prs.slide_width = Inches(8.26)
    prs.slide_height = Inches(11.69)
    return prs, prs.slide_layouts[6]


def _use_year_label(data_rows):
    """前25年里非年龄标记行>=10个 → 按年份显示，否则按岁数显示"""
    yr_set = {r["year"] for r in data_rows if not r.get("is_age_marker")}
    return sum(1 for y in range(1, 26) if y in yr_set) >= 10


def _ylab(r, last, use_year=True):
    if r.get("is_age_marker"):
        # 年龄标记行：若年份可由年龄推算（year 存在），统一显示"第X年"，与其余行一致
        if r.get("year"):
            return f"第 {r['year']} 年"
        return f"{r['age']}岁"
    return f"第 {r['year']} 年"


def _premium_desc(b, cur):
    """Slide1 缴费描述富文本段：普通部分 SUB，金额/年数部分红加粗。"""
    pt = b.get("premium_type") or ""
    annual = b.get("annual_premium") or 0
    total = b.get("total_premium") or annual
    years = b.get("pay_years") or 0
    if ("趸" in pt) or (years and years <= 1):
        return [("趸缴 ", SUB, False), (_wan(annual), RED, True), (" " + cur, SUB, False)]
    return [("年缴 ", SUB, False), (_wan(annual), RED, True), (f" {cur}，缴 ", SUB, False),
            (str(years), RED, True), (" 年，合共 ", SUB, False),
            (_wan(total), RED, True), (" " + cur, SUB, False)]


def add_base_slides(prs, blank, data, rate=6.8):
    """投入方案 + 正常/乐观/保证 三张基础表。"""
    b = data["basic"]
    cl = _cl(data)
    cur = b.get("currency", "美元")
    rows = pick_rows(data["rows"])
    last = rows[-1] if rows else None
    uy = _use_year_label(data["rows"])
    ylab = lambda r: _ylab(r, last, uy)
    age = b.get("anb_age", "")

    # === Slide 1：投入方案 + 正常预期退保价值 ===
    slide = prs.slides.add_slide(blank)
    _box(slide, data["product_name"], 0.71, 0.15, 6.85, size=24, bold=True, color=RED, align=PP_ALIGN.CENTER)
    _richbox(slide, 0.71, 0.69, 6.85, 16, _premium_desc(b, cur), align=PP_ALIGN.CENTER)
    _box(slide, f"{age}岁｜保单货币：{cur}", 0.71, 1.14, 6.85, size=12, color=GRAY, align=PP_ALIGN.CENTER)
    _add_line(slide, 0.52, 1.54, 7.08)
    _box(slide, "退保价值", 0.52, 1.60, 7.22, size=20, bold=True, color=DARK, align=PP_ALIGN.CENTER)
    normal_rows = [[ylab(r), _wan(r["surrender"]), _rmb(r["surrender"], rate)] for r in rows]
    _build_table(slide, 2.05, [["保单年度", cur, f"{cl}（约）"]], normal_rows,
                 col_w=[2.62, 2.3, 2.3], row_h=0.42, hdr_size=24, data_size=18, red_cols=[2])
    _box(slide, "剩余资金可一次性取出，或传承给下一代", 0.52, 10.95, 7.22, size=16, color=RED, align=PP_ALIGN.CENTER)
    _footer(slide, rate, cl)

    # === Slide 2：乐观预期 ===
    has_opt = bool(data.get("optimistic"))
    def opt_of(r):
        o = data["optimistic"].get(r["year"]) or data["optimistic"].get(str(r["year"])) or {}
        return o.get("optimistic", r["surrender"])
    if has_opt:
        slide = prs.slides.add_slide(blank)
        _box(slide, "退保价值", 3.43, 0.20, 1.44, size=22, bold=True, color=DARK, align=PP_ALIGN.CENTER)
        _box(slide, "（投资回报比现时预期每年上升约 1.4%）", 0.52, 0.72, 7.22, size=16, color=DARK, align=PP_ALIGN.CENTER)
        opt_rows = [[ylab(r), _wan(opt_of(r)), _rmb(opt_of(r), rate)] for r in rows]
        _build_table(slide, 1.14, [["保单年度", cur, f"{cl}（约）"]], opt_rows,
                     col_w=[2.62, 2.3, 2.3], row_h=0.45, hdr_size=24, data_size=18, red_cols=[2])
        _box(slide, "长远受每年 6.5% 客户总内部回报率规限，乐观与正常情景趋于一致",
             0.52, 11.0, 7.22, size=16, color=RED, align=PP_ALIGN.CENTER)
        _box(slide, "乐观预期", 6.93, 10.76, 0.97, size=12, color=INK0, align=PP_ALIGN.LEFT)
        _footer(slide, rate, cl)

    # === Slide 3：保证 vs 非保证 ===
    if has_opt:
        slide = prs.slides.add_slide(blank)
        _box(slide, "保证收益 vs 非保证收益", 0.0, 0.20, 8.26, size=22, bold=True, color=DARK, align=PP_ALIGN.CENTER)
        _box(slide, "保证现金价值（无论市场表现如何，这部分金额确定可获得）",
             0.52, 0.85, 7.22, size=16, color=GRAY, align=PP_ALIGN.CENTER)
        g_rows = [[ylab(r), _wan(r["guaranteed"]), _rmb(r["guaranteed"], rate),
                   _wan(r["bonus"]), _rmb(r["bonus"], rate)] for r in rows]
        _build_table(slide, 1.25,
                     [["保单年度", "保证收益", "", "非保证收益", ""],
                      ["", cur, f"{cl}（约）", cur, f"{cl}（约）"]],
                     g_rows, col_w=[1.34, 1.41, 1.24, 1.52, 1.71], row_h=0.44,
                     two_row_header=True, hdr_size=18, hdr2_size=14, data_size=14,
                     zebra=False, red_cols=[2, 4])
        _box(slide, "保证收益为保险公司承诺的最低现金价值，不包含非保证红利与终期分红",
             0.43, 11.0, 7.38, size=16, color=GRAY, align=PP_ALIGN.CENTER)
        _footer(slide, rate, cl)


def _para_box(slide, l, t, w, lines, size=12, color=DARK, align=PP_ALIGN.LEFT, h=0.5):
    """多行文本框。"""
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = "微软雅黑"
    return box


def _est_lines(s, per_line=34):
    """估算一段中文文本在 7 英寸文本框内的换行行数，用于纵向排版递进。"""
    if not s:
        return 1
    return max(1, (len(s) + per_line - 1) // per_line)


def add_cim3_slides(prs, blank, data, rate=6.8):
    """CIM3（诚保一生危疾保险计划）：基础信息 + 基本保障 + 退保价值(单利复利+币种换算) + 产品优势。"""
    b = data["basic"]
    cl = _cl(data)
    cur = b.get("currency", "美元")
    all_rows = data["rows"]
    rows = pick_rows(all_rows)
    last_y = rows[-1]["year"] if rows else None
    has_gap = any(r.get("is_age_marker") and r["year"] != last_y for r in rows)
    age = b.get("anb_age", "")
    if has_gap:
        ylab = lambda r: (f"{r['age']}岁") if r.get("is_age_marker") else f"第 {r['year']} 年"
    else:
        ylab = lambda r: f"第 {r['year']} 年"

    # === Slide 1：基础信息 ===
    slide = prs.slides.add_slide(blank)
    _box(slide, data["product_name"], 0.71, 0.15, 6.85, size=24, bold=True, color=RED, align=PP_ALIGN.CENTER)
    _richbox(slide, 0.71, 0.69, 6.85, 16, _premium_desc(b, cur), align=PP_ALIGN.CENTER)
    _box(slide, f"{age}岁｜保单货币：{cur}", 0.71, 1.14, 6.85, size=12, color=GRAY, align=PP_ALIGN.CENTER)
    _add_line(slide, 0.52, 1.54, 7.08)
    _box(slide, "基础信息", 0.52, 1.60, 7.22, size=20, bold=True, color=DARK, align=PP_ALIGN.CENTER)
    base_rows = [
        ["基本保额", _wan(b.get("basic_amount") or 0), _rmb(b.get("basic_amount") or 0, rate)],
        ["首10年赠送保障（升级保障）", _wan(b.get("gift10") or 0), _rmb(b.get("gift10") or 0, rate)],
        ["总保障额", _wan(b.get("total_amount") or 0), _rmb(b.get("total_amount") or 0, rate)],
        ["保障期", b.get("coverage") or "终身", ""],
        ["供款期", f"{b.get('pay_years')} 年" if b.get("pay_years") else "-", ""],
        ["年保费", f"{int(b.get('annual_premium') or 0):,}", f"{int((b.get('annual_premium') or 0) * rate):,}"],
        ["总保费", f"{int(b.get('total_premium') or 0):,}", f"{int((b.get('total_premium') or 0) * rate):,}"],
    ]
    _build_table(slide, 2.05, [["项目", cur, f"{cl}（约）"]], base_rows,
                 col_w=[3.2, 2.0, 2.0], row_h=0.40, hdr_size=24, data_size=15, red_cols=[2])
    _box(slide, "基本保障 · 退保价值 · 产品优势 详见后页", 0.52, 10.9, 7.22, size=16, color=RED, align=PP_ALIGN.CENTER)
    _footer(slide, rate, cl)

    gift_text = f"{_wan(b.get('basic_amount') or 0)}×9次"
    ann_text = "6,000美元/年"
    life_text = f"{_wan(b.get('basic_amount') or 0)}美元"

    # === Slide 2：基本保障 ===
    slide = prs.slides.add_slide(blank)
    _box(slide, "基本保障", 3.11, 0.20, 2.0, size=22, bold=True, color=DARK, align=PP_ALIGN.CENTER)
    _box(slide, "再次重疾每次100%基本保额、最多9次；年金每年6,000美元；人寿延伸保障100,000美元（86岁前身故）",
         0.52, 0.72, 7.22, size=12, color=GRAY, align=PP_ALIGN.CENTER)
    prot_rows = [[ylab(r), _wan(r["paid"]), _wan(r["death_total"]),
                  gift_text, ann_text, life_text, _wan(r["surrender"])] for r in rows]
    _build_table(slide, 1.25,
                 [["保单年度", "已缴保费", "首次重疾或身故", "再次重疾(×9)", "严重脑退化及帕金森年金", "人寿延伸保障", "退保价值"]],
                 prot_rows, col_w=[1.05, 1.0, 1.12, 1.02, 1.1, 1.02, 0.93], row_h=0.42,
                 hdr_size=13, data_size=12, red_cols=[6])
    _footer(slide, rate, cl)

    # === Slide 3：退保价值与回报率 ===
    slide = prs.slides.add_slide(blank)
    _box(slide, "退保价值与回报率", 2.35, 0.20, 3.6, size=22, bold=True, color=DARK, align=PP_ALIGN.CENTER)
    _box(slide, "单利/复利为基于总保费的估算年化回报；保单前期退保价值低于已缴保费，故回报率以 / 显示",
         0.52, 0.72, 7.22, size=12, color=GRAY, align=PP_ALIGN.CENTER)
    sr_rows = []
    for r in rows:
        s = f"{r['simple_pct']}%" if (r.get("simple_pct") is not None and r["simple_pct"] > 0) else "/"
        c = f"{r['compound_pct']}%" if (r.get("compound_pct") is not None and r["compound_pct"] > 0) else "/"
        sr_rows.append([ylab(r), _wan(r["paid"]), _wan(r["death_total"]),
                        _wan(r["surrender"]), _rmb(r["surrender"], rate), s, c])
    _build_table(slide, 1.25,
                 [["保单年度", "已缴保费", "首次重疾或身故", f"退保价值({cur})", f"{cl}（约）", "单利%/年", "复利%/年"]],
                 sr_rows, col_w=[1.05, 1.0, 1.12, 1.1, 1.05, 0.95, 0.95], row_h=0.42,
                 hdr_size=13, data_size=12, red_cols=[4])
    _footer(slide, rate, cl)

    # === Slide 4：产品优势（保障明细表）===
    slide = prs.slides.add_slide(blank)
    _box(slide, "产品优势", 3.2, 0.18, 1.9, size=22, bold=True, color=DARK, align=PP_ALIGN.CENTER)
    _box(slide, f"美元汇率：{rate}（仅供参考）", 0.55, 0.60, 7.2, size=10, color=GRAY)
    G = b.get("basic_amount") or 0
    gift = b.get("gift10") or 0
    groups = [
        ("危疾保障", [
            ("第一次重疾理赔", "覆盖127种病况；首10年升级保障总额150%", G + gift),
            ("重症多次理赔", "延伸保障最高累计900%保额、每笔100%（最多9次）", G * 9),
        ]),
        ("特殊关怀", [
            ("深切治疗ICU保障", "连续入住ICU三日或以上，赔基本保额10%", G * 0.1),
            ("脑退化/帕金森病终身年金", "确诊后每年支付基本保额6%，直至身故", f"{int(G * 0.06):,}/年"),
            ("儿童专属保障", "严重湿疹、荨麻疹、肺炎住院津贴等", None),
        ]),
        ("人寿与身故", [
            ("人寿延伸保障", "严重疾病理赔后、86岁前身故，额外赔付", G),
        ]),
        ("保费与退保", [
            ("保费豁免", "确诊严重疾病后，豁免日后所有保费", None),
            ("退保提取", "连本带利取回保单现金价值", None),
        ]),
    ]
    n_rows = 1 + sum(1 + len(rows) for _, rows in groups)
    tbl = slide.shapes.add_table(n_rows, 3, Inches(0.52), Inches(0.98), Inches(7.22), Inches(0.5 * n_rows)).table
    for i, w in enumerate([3.9, 1.77, 1.55]):
        tbl.columns[i].width = Inches(w)
    tbl.rows[0].height = Inches(0.32)
    for j, h in enumerate(["保障内容", "美元", "人民币（约）"]):
        _style_cell(tbl.cell(0, j), h, bold=True, color=WHITE, size=11, fill=RED)
    r = 1
    for gi, (g, rows) in enumerate(groups):
        _style_cell_lines(tbl.cell(r, 0), [(g, True, RED, 11)], align=PP_ALIGN.LEFT, fill=LIGHT)
        r += 1
        for name, note, v in rows:
            _style_cell_lines(tbl.cell(r, 0),
                              [(name, True, DARK, 10), (note, False, GRAY, 8)],
                              align=PP_ALIGN.LEFT, fill=None)
            usd = v if isinstance(v, str) else (_wan(v) if v is not None else "—")
            cny = "—" if (v is None or isinstance(v, str)) else _wan(v * rate)
            _style_cell(tbl.cell(r, 1), usd, size=10, color=RED)
            _style_cell(tbl.cell(r, 2), cny, size=10)
            tbl.rows[r].height = Inches(0.5)
            r += 1
    _footer(slide, rate, cl)

    # === Slide 5：产品优势（文字版 1/2）—— 生/病/老/死 分块 ===
    advs = data.get("product_advantage") or []
    slide = prs.slides.add_slide(blank)
    _box(slide, "产品优势", 3.2, 0.18, 1.9, size=22, bold=True, color=DARK, align=PP_ALIGN.CENTER)
    top = 0.70
    if advs:
        a0 = advs[0]
        _box(slide, "■ " + a0["title"], 0.55, top, 7.2, size=15, color=RED, bold=True)
        top += 0.40
        for it in a0["items"]:
            parts = it.split("·", 1)
            tag = parts[0]
            body = parts[1] if len(parts) > 1 else ""
            _richbox_red(slide, 0.55, top, 7.15, 12, [(f"{tag}  ", RED, True), (body, DARK, False)], h=0.9)
            top += 0.04 + 0.26 * _est_lines(body)
    _footer(slide, rate, cl)

    # === Slide 6：产品优势（文字版 2/2）—— 优势 2/3 ===
    slide = prs.slides.add_slide(blank)
    _box(slide, "产品优势", 3.2, 0.18, 1.9, size=22, bold=True, color=DARK, align=PP_ALIGN.CENTER)
    top = 0.70
    for a in advs[1:]:
        _box(slide, "■ " + a["title"], 0.55, top, 7.2, size=15, color=RED, bold=True)
        top += 0.40
        for it in a["items"]:
            _richbox_red(slide, 0.78, top, 7.0, 12, [(it, DARK, False)], h=0.3 * _est_lines(it) + 0.1)
            top += 0.10 + 0.26 * _est_lines(it)
        top += 0.14
    _footer(slide, rate, cl)


def add_withdrawal_slide(prs, blank, data, rate=6.8, label=None):
    """提取方案页。label 用于多方案序号（一/二/三…）。"""
    w = data.get("withdrawal")
    if not w or not w.get("start_year"):
        return
    cl = _cl(data)
    cur = data["basic"].get("currency", "美元")
    slide = prs.slides.add_slide(blank)
    wrows = [x for x in w["rows"] if x["withdrawal"] > 0]
    n_years = len(wrows)
    wlast = wrows[-1]
    annual = w["annual_withdrawal"]
    total = wlast["accumulated"]

    _box(slide, "建议提取方案" + (label or ""), 3.05, 0.30, 2.17, size=20, bold=True, color=INK0, align=PP_ALIGN.CENTER)
    _box(slide, f"从第 {w['start_year']} 年开始提取，共提取 {n_years} 年", 2.72, 0.80, 2.81, size=12, color=DARK, align=PP_ALIGN.CENTER)
    _richbox(slide, 2.54, 1.15, 3.18, 20,
             [("每年领取 ", RED, True), (f"{annual:,.0f} ", RED, True), (cur, RED, True)],
             align=PP_ALIGN.CENTER)
    _box(slide, f"（每年约 {_rmb(annual, rate)}{cl}，每月约 {_rmb(annual/12, rate)}{cl}）",
         1.92, 1.55, 4.13, size=12, color=DARK, align=PP_ALIGN.CENTER)
    _richbox(slide, 1.32, 1.9, 5.63, 16,
             [(f"{n_years} 年总共领取 {_wan(total)}{cur}（约 {_rmb(total, rate)}{cl}）", DARK, True)],
             align=PP_ALIGN.CENTER)
    _box(slide, "提取后账户剩余的总现金价值", 2.21, 2.40, 3.85, size=20, bold=True, color=DARK, align=PP_ALIGN.CENTER)
    _add_line(slide, 0.59, 2.33, 7.08)
    key = [x for x in wrows if x["year"] == w["start_year"] or x["year"] % 5 == 0 or x == wrows[-1]]
    rem_rows = [[f"第 {x['year']} 年",
                 _wan(x["surrender"]), _rmb(x["surrender"], rate)] for x in key]
    _build_table(slide, 2.85, [["保单年度", cur + "（剩余价值）", f"{cl}（约）"]], rem_rows,
                 col_w=[2.62, 2.3, 2.3], row_h=0.40, hdr_size=18, data_size=14, red_cols=[2])
    _richbox(slide, 0.59, 11.0, 7.15, 16,
             [("领完一辈子，账户还剩近 ", DARK, False),
              (_wan(wlast["surrender"]) + " ", RED, False),
              ("继续传承给下一代", DARK, False)],
             align=PP_ALIGN.CENTER)
    _footer(slide, rate, cl)


def _save(prs):
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def build_pptx(data, rate=6.8):
    """单份：基础表 + 提取页（CIM3 走专属页）。"""
    prs, blank = _new_prs()
    if data.get("product_code") == "CIM3":
        add_cim3_slides(prs, blank, data, rate)
    else:
        add_base_slides(prs, blank, data, rate)
        add_withdrawal_slide(prs, blank, data, rate)
    return _save(prs)


def build_combined_pptx(base, variants, rate=6.8):
    """同一投入方案多提取方案：基础表一次，每个提取方案各一页。"""
    cn = "一二三四五六七八九十"
    prs, blank = _new_prs()
    add_base_slides(prs, blank, base, rate)
    for i, v in enumerate(variants):
        label = cn[i] if i < len(cn) else str(i + 1)
        add_withdrawal_slide(prs, blank, v, rate, label=label)
    return _save(prs)
